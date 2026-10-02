"""Build the Red Flag Ritual brand deck — minimalist, old money.

Build Mode with the pptx-designer public API. All imagery is bespoke (make_art.py) — no stock photography.
Fonts: Libre Caslon Display + EB Garamond (see fonts/). Run: python make_art.py && python build_deck.py
"""
from pathlib import Path

from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from pptx_designer import Presentation
from pptx_designer.tools.layout import add_slide, clean_save
from pptx_designer.tools.shapes import rect
from pptx_designer.tools.text import text

import motion

HERE = Path(__file__).resolve().parent
ART = HERE / "art"
OUT = HERE / "output"

CASLON, SERIF = "Libre Caslon Display", "EB Garamond"
W, H = 13.333, 7.5
TOTAL = 10
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]

PAPERS = {
    "ivory": ("paper_ivory", {"ink": "#1B1A17", "body": "#4A4439", "muted": "#8C8270", "rule": "#B9A27A"}),
    "stone": ("paper_stone", {"ink": "#1B1A17", "body": "#3F392F", "muted": "#7D7362", "rule": "#A88A57"}),
    "ink": ("paper_ink", {"ink": "#F4EFE4", "body": "#D8D0C0", "muted": "#9A9080", "rule": "#A88A57"}),
    "oxblood": ("paper_oxblood", {"ink": "#F4EFE4", "body": "#E6D9CC", "muted": "#C9AFA4", "rule": "#C9A66B"}),
}
ACC = {"oxblood": "#5E1A1D", "green": "#24382B", "brass": "#A88A57", "ivory": "#F4EFE4", "paper": "#ECE5D7"}


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.pal = PAPERS["ivory"][1]
        self.plan = {}

    def col(self, k):
        return k if k.startswith("#") else (self.pal.get(k) or ACC[k])

    @property
    def C(self):
        return {"background": "#F4EFE4", "text_body": self.pal["body"], "text_dark": self.pal["ink"],
                "text_muted": self.pal["muted"], "font_body": SERIF, "font_heading": CASLON}

    def T(self, s, x, y, w, h, txt, size=14, font=SERIF, color="body", italic=False, align="left",
          spacing=None, leading=None, caps=False, bold=False):
        box = text(s, x, y, w, h, txt.upper() if caps else txt, font_size=size, color=self.col(color), bold=bold,
                   align=align, font_name=font, C=self.C)
        tf = box.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for p in tf.paragraphs:
            p.alignment = tf.paragraphs[0].alignment
            if leading:
                p.line_spacing = leading
            for r in p.runs:
                r.font.italic = italic
                if spacing is not None:
                    r.font._element.set("spc", str(spacing))
        return box

    def H1(self, s, x, y, w, h, txt, size=44, color="ink", align="left", leading=1.0, italic=False):
        return self.T(s, x, y, w, h, txt, size=size, font=CASLON, color=color, align=align, leading=leading,
                      italic=italic)

    def small_caps(self, s, x, y, w, txt, color="muted", size=8.5, align="left", spacing=400):
        return self.T(s, x, y, w, 0.25, txt, size=size, color=color, caps=True, spacing=spacing, align=align)

    def rule(self, s, x, y, w, color="rule", t=0.01):
        return rect(s, x, y, w, t, self.col(color), C=self.C)

    def vrule(self, s, x, y, h, color="rule"):
        return rect(s, x, y, 0.01, h, self.col(color), C=self.C)

    def frame(self, s, inset=0.32, gap=0.06, color="rule"):
        """Stationery double rule around the page."""
        for d in (inset, inset + gap):
            r = rect(s, d, d, W - 2 * d, H - 2 * d, "#000000", line=self.col(color), C=self.C)
            r.fill.background()
            r.line.width = Pt(0.6)
            r.name = "static-frame"

    def art(self, s, name, x, y, w=None, h=None):
        from PIL import Image
        iw, ih = Image.open(ART / f"{name}.png").size
        if w and not h:
            h = w * ih / iw
        elif h and not w:
            w = h * iw / ih
        pic = s.shapes.add_picture(str(ART / f"{name}.png"), Inches(x), Inches(y), Inches(w), Inches(h))
        pic.name = "static-art"
        return pic

    def slide(self, paper="ivory", transition=("fade", False), folio=True, frame=True):
        name, self.pal = PAPERS[paper]
        s = add_slide(self.prs)
        n = len(self.prs.slides)
        self.plan[n] = transition
        bg = s.shapes.add_picture(str(ART / f"{name}.jpg"), 0, 0, Inches(W), Inches(H))
        bg.name = "static-bg"
        if frame:
            self.frame(s)
        if folio:
            for shp in (self.small_caps(s, 0, H - 0.72, W, "Red Flag Ritual", align="center", size=7.5),
                        self.T(s, 0, H - 0.98, W, 0.25, ROMAN[n - 1], size=10, font=CASLON, color="rule",
                               align="center")):
                shp.name = "chrome"
        return s

    def flagmark(self, s, x, y, h=0.42, ivory=False):
        return self.art(s, "flag_ivory" if ivory else "flag", x, y, h=h)

    def save(self):
        for i, sl in enumerate(self.prs.slides, start=1):
            kind, black = self.plan.get(i, ("fade", False))
            motion.transition(sl, kind, through_black=black)
            motion.choreograph(sl, step_ms=140, budget_ms=2200)
        OUT.mkdir(exist_ok=True)
        path = OUT / "Red_Flag_Ritual.pptx"
        clean_save(self.prs, str(path))
        print(path)


def build():
    d = Deck()
    T, H1, sc = d.T, d.H1, d.small_caps

    # I — Cover ---------------------------------------------------------------------------------------
    s = d.slide(transition=("fade", True), folio=False)
    d.flagmark(s, W / 2 - 0.13, 0.75, h=0.5)
    sc(s, 0, 1.42, W, "Brand presentation", align="center", spacing=600)
    H1(s, 0, 1.75, W, 1.1, "Red Flag Ritual", size=66, align="center")
    T(s, 0, 2.95, W, 0.5, "Green flags only.", size=22, italic=True, color="ink", align="center")
    d.rule(s, W / 2 - 0.6, 3.55, 1.2)
    T(s, 0, 3.7, W, 0.4, "Signature-scent amenities for stays worth remembering.", size=13, color="body",
      align="center")
    d.art(s, "heirlooms", W / 2 - 2.45, 4.15, w=4.9)

    # II — Guests forget -------------------------------------------------------------------------------
    s = d.slide()
    sc(s, 1.1, 1.3, 6, "The problem")
    H1(s, 1.1, 1.65, 6.2, 1.8, "Guests forget\nmost stays.", size=50, leading=1.0)
    d.rule(s, 1.1, 3.55, 0.8, "oxblood")
    T(s, 1.1, 3.8, 5.4, 2.2, "Generic sachets. Mismatched bottles. A room that smells like nothing.\n\n"
      "Guests don’t complain about it. They just don’t come back — and they don’t mention you in reviews.",
      size=15, leading=1.35)
    d.art(s, "sachet", 7.35, 1.5, w=4.9)
    T(s, 7.35, 4.55, 4.9, 0.3, "The stay nobody remembers.", size=11, italic=True, color="muted", align="center")

    # III — Meet the ritual ---------------------------------------------------------------------------
    s = d.slide()
    d.art(s, "shelf", 0.85, 1.2, w=6.6)
    sc(s, 8.05, 1.35, 4.5, "The answer")
    H1(s, 8.05, 1.7, 4.6, 0.9, "Meet the ritual.", size=42)
    d.rule(s, 8.05, 2.6, 0.8, "green")
    T(s, 8.05, 2.85, 4.4, 3.4, "Red Flag Ritual turns your bathroom and bedroom into one quiet, luxurious experience.\n\n"
      "One signature scent runs through every bottle, every spray and every guest kit. It arrives at your door every "
      "month. You never think about restocking again.", size=14, leading=1.4)

    # IV — No. 01 Late Checkout (stone) ----------------------------------------------------------------
    s = d.slide(paper="stone", transition=("fade", True))
    d.art(s, "mist", 5.4, 0.85, w=7.4)
    sc(s, 1.1, 1.3, 5, "Our signature scent", color="oxblood")
    H1(s, 1.1, 1.65, 5, 0.9, "No. 01", size=30, color="muted")
    H1(s, 1.1, 2.2, 5.5, 1.0, "Late Checkout", size=52, italic=True)
    T(s, 1.1, 3.35, 4.4, 1.0, "Warm, clean, slow-morning. It smells like you didn’t have to leave yet.", size=15,
      leading=1.35)
    d.rule(s, 1.1, 4.6, 4.3)
    T(s, 1.1, 4.8, 4.3, 1.2, "“Guests remember how a place made them feel. Scent is how they remember it.”",
      size=17, italic=True, color="ink", leading=1.3)

    # V — The collection ------------------------------------------------------------------------------
    s = d.slide()
    sc(s, 0, 0.95, W, "The collection", align="center", spacing=600)
    H1(s, 0, 1.22, W, 0.8, "Six bottles. One scent.", size=34, align="center")
    coll = [("c_sud", "SUD", "Shampoo", "Rich, soft lather. Leaves hair clean without stripping it.", False),
            ("c_dew", "DEW", "Body wash", "Fresh, light wash. The scent stays on the skin after the shower.", False),
            ("c_vel", "VEL", "Body lotion", "Velvet-soft, fast-absorbing, non-greasy finish.", False),
            ("c_slk", "SLK", "Conditioner", "Smooth, silky detangle.", True),
            ("c_air", "AIR", "Room spray", "Signature scent for rooms and linen. The first thing guests notice.", False),
            ("c_lng", "LNG", "Diffuser refill", "Short for “linger”. Keeps the scent in the room all day.", True)]
    cw = (W - 2.0) / 6
    for i, (img, code, name, what, estate) in enumerate(coll):
        x = 1.0 + i * cw
        d.art(s, img, x + cw / 2 - 0.82, 1.95, w=1.64)
        T(s, x, 4.72, cw, 0.4, code, size=24, font=CASLON, color="ink", align="center", spacing=600)
        T(s, x, 5.15, cw, 0.3, name, size=12, italic=True, color="ink", align="center")
        T(s, x + 0.1, 5.45, cw - 0.2, 0.8, what, size=9.5, color="body", align="center", leading=1.25)
        if estate:
            sc(s, x, 6.2, cw, "Estate only", color="oxblood", size=7, align="center", spacing=300)

    # VI — Choose your house rules -----------------------------------------------------------------------
    s = d.slide()
    sc(s, 0, 0.95, W, "Plans", align="center", spacing=600)
    H1(s, 0, 1.22, W, 0.8, "Choose your house rules.", size=36, align="center")
    T(s, 0, 1.92, W, 0.3, "Every plan includes 30 guest sets a month, delivered to your property.  ·  "
      "Prices exclusive of applicable taxes. Cancel anytime.", size=12,
      italic=True, color="body", align="center")
    tiers = [("House", "₹999", "≈ ₹33 per guest", "tier_house",
              ["SUD + DEW in 10 mL vials", "AIR room spray"], False),
             ("Manor", "₹1,999", "≈ ₹67 per guest", "tier_manor",
              ["SUD + DEW + VEL in 10 mL amber glass", "AIR glass spray"], True),
             ("Estate", "₹3,999", "≈ ₹133 per guest", "tier_estate",
              ["SUD + DEW + VEL + SLK in 15 mL designer glass", "AIR glass spray", "LNG diffuser"], False)]
    tw = 3.55
    x0 = (W - 3 * tw - 2 * 0.3) / 2
    for i, (name, price, per, img, inc, hi) in enumerate(tiers):
        x = x0 + i * (tw + 0.3)
        top = 2.45
        if hi:
            panel = rect(s, x, top, tw, 3.9, ACC["green"], C=d.C)
            panel.name = "static-panel"
            fg, mut = "#F4EFE4", "#C9C1B0"
            sc(s, x, top + 0.18, tw, "Most hosts", color="#C9A66B", size=7.5, align="center", spacing=500)
        else:
            fg, mut = d.pal["ink"], d.pal["muted"]
            d.rule(s, x, top, tw)
            d.rule(s, x, top + 3.9, tw)
        d.art(s, img, x + 0.25, top + 0.5, h=2.05)
        H1(s, x + 1.55, top + 0.55, 2.0, 0.5, name, size=24, color=fg)
        H1(s, x + 1.55, top + 1.05, 2.0, 0.6, price, size=30, color=fg)
        T(s, x + 1.55, top + 1.62, 2.0, 0.3, "per month", size=10, italic=True, color=mut)
        T(s, x + 1.55, top + 1.9, 2.0, 0.3, per, size=10, color=mut)
        for j, line in enumerate(inc):
            T(s, x + 0.3, top + 2.75 + j * 0.42, tw - 0.6, 0.4, line, size=11, color=fg, align="center")

    # VII — Ritual Classic (ink) -------------------------------------------------------------------------
    s = d.slide(paper="ink", transition=("fade", True))
    d.art(s, "kit", 6.2, 0.95, w=6.4)
    sc(s, 1.1, 1.3, 5, "The guest take-home kit", color="brass")
    H1(s, 1.1, 1.65, 5, 0.9, "Ritual Classic", size=48)
    d.rule(s, 1.1, 2.6, 0.8)
    T(s, 1.1, 2.85, 4.6, 2.0, "Four test tubes of the ritual in a monogrammed box: SUD, DEW, VEL and AIR.\n\n"
      "A thank-you guests keep — and a QR code that lets them bring the scent home.", size=14, leading=1.4)
    H1(s, 1.1, 4.95, 4, 0.7, "₹149", size=40, color="#C9A66B")
    T(s, 2.55, 5.12, 3.5, 0.5, "per kit  ·  minimum 10\nadd to any plan", size=11, italic=True, leading=1.2)

    # VIII — How it works ------------------------------------------------------------------------------
    s = d.slide()
    sc(s, 0, 1.25, W, "How it works", align="center", spacing=600)
    H1(s, 0, 1.52, W, 0.8, "Set it once. It simply repeats.", size=34, align="center")
    steps = [("i_plan", "Pick your plan.", "House, Manor or Estate."),
             ("i_deliver", "We deliver monthly.", "30 guest sets, sealed and ready."),
             ("i_notice", "Guests notice.", "The scent shows up in reviews."),
             ("i_repeat", "It just repeats.", "Auto-renews every month. Change or cancel anytime.")]
    sw = (W - 2.4) / 4
    d.rule(s, 1.2 + sw / 2, 3.5, sw * 3)
    for i, (ic, t, b) in enumerate(steps):
        x = 1.2 + i * sw
        circ = s.shapes.add_shape(9, Inches(x + sw / 2 - 0.6), Inches(2.9), Inches(1.2), Inches(1.2))
        circ.fill.solid()
        circ.fill.fore_color.rgb = __import__("pptx").dml.color.RGBColor.from_string("F4EFE4")
        circ.line.color.rgb = __import__("pptx").dml.color.RGBColor.from_string("A88A57")
        circ.line.width = Pt(0.75)
        d.art(s, ic, x + sw / 2 - 0.38, 3.12, w=0.76)
        T(s, x, 4.35, sw, 0.4, ROMAN[i], size=14, font=CASLON, color="rule", align="center")
        H1(s, x + 0.1, 4.75, sw - 0.2, 0.5, t, size=20, align="center")
        T(s, x + 0.25, 5.3, sw - 0.5, 0.9, b, size=12, align="center", leading=1.3)

    # IX — Why hosts choose ----------------------------------------------------------------------------
    s = d.slide()
    sc(s, 1.1, 1.3, 6, "Why hosts choose Red Flag Ritual")
    H1(s, 1.1, 1.65, 6, 0.9, "Green flags, only.", size=40)
    reasons = ["Hotel-grade experience from ₹33 per guest", "One signature scent across the whole stay",
               "Sealed, single-stay bottles for every guest", "No restocking runs. It arrives every month.",
               "Trusted in Red Flag Homes properties first"]
    for i, r in enumerate(reasons):
        y = 2.75 + i * 0.66
        d.rule(s, 1.1, y, 5.8)
        T(s, 1.1, y + 0.17, 0.5, 0.35, ROMAN[i], size=12, font=CASLON, color="rule")
        T(s, 1.7, y + 0.15, 5.2, 0.4, r, size=15, color="ink")
    d.rule(s, 1.1, 2.75 + 5 * 0.66, 5.8)
    cx, cy, cw2, ch = 7.9, 2.0, 4.3, 3.6                       # review card
    card = rect(s, cx, cy, cw2, ch, ACC["paper"], C=d.C)
    card.name = "static-card"
    for dd in (0.12, 0.17):
        r = rect(s, cx + dd, cy + dd, cw2 - 2 * dd, ch - 2 * dd, "#000000", line=ACC["brass"], C=d.C)
        r.fill.background()
        r.line.width = Pt(0.5)
        r.name = "static-frame"
    T(s, cx, cy + 0.45, cw2, 0.4, "★ ★ ★ ★ ★", size=13, color="brass", align="center", spacing=200)
    H1(s, cx + 0.4, cy + 1.0, cw2 - 0.8, 1.4, "“The room smelled incredible.”", size=28, italic=True,
       align="center", leading=1.1)
    d.rule(s, cx + cw2 / 2 - 0.4, cy + 2.55, 0.8)
    sc(s, cx, cy + 2.75, cw2, "Sample review", color="muted", size=7, align="center")

    # X — The only red flag is leaving (oxblood) ----------------------------------------------------------
    s = d.slide(paper="oxblood", transition=("fade", True), folio=False)
    d.flagmark(s, W / 2 - 0.13, 1.25, h=0.55, ivory=True)
    H1(s, 0, 2.15, W, 0.9, "The only red flag", size=54, align="center")
    H1(s, 0, 3.0, W, 0.9, "is leaving.", size=54, align="center", italic=True)
    T(s, 0, 4.15, W, 0.5, "Start your ritual today.", size=20, italic=True, color="ink", align="center")
    d.rule(s, W / 2 - 0.6, 4.85, 1.2)
    sc(s, 0, 5.1, W, "redflaghomes.in   ·   [WhatsApp number]   ·   [Instagram handle]", color="body", size=9,
       align="center", spacing=300)

    assert len(d.prs.slides) == TOTAL
    d.save()


if __name__ == "__main__":
    build()
